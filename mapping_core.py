from __future__ import annotations

import csv
import os
import re
import shutil
import sqlite3
import tempfile
import threading
import time
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Callable

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ProgressCallback = Callable[[str, int, int | None, str], None]

WINDOWS_INVALID_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
WINDOWS_RESERVED_NAMES = {
    "CON",
    "PRN",
    "AUX",
    "NUL",
    *(f"COM{i}" for i in range(1, 10)),
    *(f"LPT{i}" for i in range(1, 10)),
}


class MappingCancelled(RuntimeError):
    """Raised internally when the user cancels a long-running phase."""


@dataclass(frozen=True)
class MappingConfig:
    excel_path: Path
    sheet_name: str
    identifier_column: int
    filename_column: int
    source_folder: Path
    target_folder: Path
    result_path: Path
    case_sensitive: bool = True
    include_subfolders: bool = False


@dataclass(frozen=True)
class PreviewSummary:
    total_rows: int
    valid_rows: int
    distinct_requested_files: int
    source_files_scanned: int
    matched_rows: int
    matched_source_files: int
    not_found_rows: int
    conflict_rows: int
    invalid_rows: int
    folders_to_create: int
    same_volume: bool
    elapsed_seconds: float


@dataclass(frozen=True)
class ExecutionSummary:
    files_moved: int
    rows_mapped: int
    rows_failed: int
    rows_not_processed: int
    result_path: Path
    log_path: Path
    cancelled: bool
    elapsed_seconds: float


class TaskControl:
    def __init__(self) -> None:
        self._cancel = threading.Event()
        self._pause = threading.Event()

    def cancel(self) -> None:
        self._cancel.set()
        self._pause.clear()

    def toggle_pause(self) -> bool:
        if self._pause.is_set():
            self._pause.clear()
            return False
        self._pause.set()
        return True

    @property
    def cancelled(self) -> bool:
        return self._cancel.is_set()

    @property
    def paused(self) -> bool:
        return self._pause.is_set()

    def checkpoint(self) -> None:
        while self._pause.is_set() and not self._cancel.is_set():
            time.sleep(0.1)
        if self._cancel.is_set():
            raise MappingCancelled()


def _as_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def normalize_filename(filename: str, case_sensitive: bool) -> str:
    # Compatibility argument retained; matching is always exact in v1.1.0.
    return filename


def validate_identifier(value: str) -> str | None:
    if not value:
        return "Identifier is blank"
    if WINDOWS_INVALID_CHARS.search(value):
        return "Identifier contains characters that Windows does not allow in folder names"
    if value.endswith((" ", ".")):
        return "Identifier ends with a space or period, which Windows does not allow"
    if value.split(".", 1)[0].upper() in WINDOWS_RESERVED_NAMES:
        return "Identifier is a reserved Windows folder name"
    return None


def validate_filename(value: str) -> str | None:
    if not value:
        return "File name is blank"
    if Path(value).name != value or "/" in value or "\\" in value:
        return "File name must not contain a folder path"
    if WINDOWS_INVALID_CHARS.search(value):
        return "File name contains characters that Windows does not allow"
    if value.endswith((" ", ".")):
        return "File name ends with a space or period, which Windows does not allow"
    return None


def unique_output_headers(existing: list[object]) -> tuple[str, str]:
    names = {_as_text(value).casefold() for value in existing if _as_text(value)}

    def choose(base: str) -> str:
        candidate = base
        counter = 2
        while candidate.casefold() in names:
            candidate = f"{base} {counter}"
            counter += 1
        names.add(candidate.casefold())
        return candidate

    return choose("Mapping Status"), choose("Failure Reason")


class MappingEngine:
    """Disk-backed mapping engine for large workbooks and file collections."""

    def __init__(
        self,
        config: MappingConfig,
        progress: ProgressCallback | None = None,
        control: TaskControl | None = None,
    ) -> None:
        self.config = config
        self.progress = progress or (lambda *_: None)
        self.control = control or TaskControl()
        self.preview_summary: PreviewSummary | None = None
        temp = tempfile.NamedTemporaryFile(prefix="file_mapping_", suffix=".sqlite", delete=False)
        temp.close()
        self.database_path = Path(temp.name)
        self._connection: sqlite3.Connection | None = None

    def close(self, keep_checkpoint: bool = False) -> None:
        if self._connection is not None:
            self._connection.close()
            self._connection = None
        if not keep_checkpoint:
            try:
                self.database_path.unlink(missing_ok=True)
            except OSError:
                pass

    def __enter__(self) -> "MappingEngine":
        return self

    def __exit__(self, *_: object) -> None:
        self.close()

    def _connect(self) -> sqlite3.Connection:
        if self._connection is None:
            connection = sqlite3.connect(self.database_path, check_same_thread=False)
            connection.execute("PRAGMA journal_mode=MEMORY")
            connection.execute("PRAGMA synchronous=OFF")
            connection.execute("PRAGMA temp_store=MEMORY")
            connection.execute("PRAGMA cache_size=-65536")
            self._connection = connection
        return self._connection

    def _reset_database(self) -> sqlite3.Connection:
        connection = self._connect()
        connection.executescript(
            """
            DROP TABLE IF EXISTS mapping_rows;
            DROP TABLE IF EXISTS target_names;
            DROP TABLE IF EXISTS source_hits;
            CREATE TABLE mapping_rows (
                row_number INTEGER PRIMARY KEY,
                identifier TEXT NOT NULL,
                file_name TEXT NOT NULL,
                norm_name TEXT NOT NULL,
                status TEXT NOT NULL,
                reason TEXT NOT NULL
            );
            CREATE TABLE target_names (
                norm_name TEXT PRIMARY KEY,
                file_name TEXT NOT NULL,
                identifier TEXT NOT NULL,
                conflict INTEGER NOT NULL DEFAULT 0
            );
            CREATE TABLE source_hits (
                norm_name TEXT NOT NULL,
                source_path TEXT NOT NULL,
                PRIMARY KEY (norm_name, source_path)
            );
            """
        )
        return connection

    def _validate_config(self) -> None:
        cfg = self.config
        if not cfg.excel_path.is_file():
            raise ValueError("The selected Excel file does not exist")
        if cfg.excel_path.suffix.casefold() != ".xlsx":
            raise ValueError("Only .xlsx files are supported")
        if not cfg.source_folder.is_dir():
            raise ValueError("The selected source folder does not exist")
        source = cfg.source_folder.resolve()
        target = cfg.target_folder.resolve()
        if source == target:
            raise ValueError("The source and target folders must be different")
        if source in target.parents:
            raise ValueError("The target folder must not be inside the source folder")
        if cfg.result_path.resolve() == cfg.excel_path.resolve():
            raise ValueError("The result Excel must not overwrite the original Excel")
        if cfg.identifier_column < 1 or cfg.filename_column < 1:
            raise ValueError("Column numbers must be one-based positive integers")
        if cfg.identifier_column == cfg.filename_column:
            raise ValueError("Identifier and file-name columns must be different")

    def precheck(self) -> PreviewSummary:
        started = time.perf_counter()
        self._validate_config()
        connection = self._reset_database()
        cfg = self.config
        self.progress("excel", 0, None, "Reading mapping rows")

        workbook = load_workbook(cfg.excel_path, read_only=True, data_only=False)
        try:
            if cfg.sheet_name not in workbook.sheetnames:
                raise ValueError(f"Worksheet not found: {cfg.sheet_name}")
            sheet = workbook[cfg.sheet_name]
            max_column = max(cfg.identifier_column, cfg.filename_column)
            excel_total = sheet.max_row - 1 if isinstance(sheet.max_row, int) else None
            batch: list[tuple[int, str, str, str, str, str]] = []
            total_rows = 0

            for row_number, cells in enumerate(
                sheet.iter_rows(min_row=2, max_col=max_column, values_only=True), start=2
            ):
                self.control.checkpoint()
                identifier = _as_text(cells[cfg.identifier_column - 1])
                file_name = _as_text(cells[cfg.filename_column - 1])
                if not identifier and not file_name:
                    continue
                total_rows += 1
                reason = validate_identifier(identifier) or validate_filename(file_name)
                status = "Move failed" if reason else "Pending"
                norm_name = normalize_filename(file_name, cfg.case_sensitive) if file_name else ""
                batch.append((row_number, identifier, file_name, norm_name, status, reason or ""))
                if len(batch) >= 5000:
                    connection.executemany("INSERT INTO mapping_rows VALUES (?, ?, ?, ?, ?, ?)", batch)
                    connection.commit()
                    batch.clear()
                    self.progress("excel", total_rows, excel_total, "Reading mapping rows")
            if batch:
                connection.executemany("INSERT INTO mapping_rows VALUES (?, ?, ?, ?, ?, ?)", batch)
                connection.commit()
        finally:
            workbook.close()

        connection.execute("CREATE INDEX idx_mapping_rows_norm ON mapping_rows(norm_name)")
        connection.execute(
            """
            INSERT INTO target_names(norm_name, file_name, identifier, conflict)
            SELECT norm_name, MIN(file_name), MIN(identifier),
                   0
            FROM mapping_rows
            WHERE status = 'Pending'
            GROUP BY norm_name
            """
        )
        connection.create_function("windows_key", 1, str.casefold)
        connection.execute("""
            UPDATE mapping_rows SET status = 'Mapping conflict',
                reason = 'Different identifiers resolve to the same Windows folder'
            WHERE windows_key(identifier) IN (
                SELECT windows_key(identifier) FROM mapping_rows
                WHERE status = 'Pending' GROUP BY windows_key(identifier)
                HAVING COUNT(DISTINCT identifier) > 1
            )
        """)
        connection.commit()

        target_count = connection.execute(
            "SELECT COUNT(*) FROM target_names WHERE conflict = 0"
        ).fetchone()[0]
        target_set: set[str] | None = None
        if target_count <= 1_000_000:
            target_set = {
                row[0]
                for row in connection.execute(
                    "SELECT norm_name FROM target_names WHERE conflict = 0"
                )
            }

        scope = "source folder and subfolders" if cfg.include_subfolders else "source folder"
        self.progress("scan", 0, None, f"Scanning the {scope}")
        lookup_cursor = connection.cursor()
        hit_batch: list[tuple[str, str]] = []
        source_files_scanned = 0
        folders: list[Path] = [cfg.source_folder]

        while folders:
            self.control.checkpoint()
            folder = folders.pop()
            try:
                entries = os.scandir(folder)
            except PermissionError as exc:
                if folder == cfg.source_folder:
                    raise ValueError(f"The source folder cannot be read: {exc}") from exc
                self.progress("scan", source_files_scanned, None, f"Skipped unreadable folder: {folder}")
                continue
            except OSError:
                continue

            with entries:
                for entry in entries:
                    self.control.checkpoint()
                    try:
                        if entry.is_file(follow_symlinks=False):
                            source_files_scanned += 1
                            norm_name = normalize_filename(entry.name, cfg.case_sensitive)
                            wanted = norm_name in target_set if target_set is not None else bool(
                                lookup_cursor.execute(
                                    "SELECT 1 FROM target_names WHERE norm_name = ? AND conflict = 0",
                                    (norm_name,),
                                ).fetchone()
                            )
                            if wanted:
                                hit_batch.append((norm_name, entry.path))
                        elif cfg.include_subfolders and entry.is_dir(follow_symlinks=False):
                            folders.append(Path(entry.path))
                        else:
                            continue
                    except OSError:
                        continue

                    if len(hit_batch) >= 1000:
                        connection.executemany("INSERT OR IGNORE INTO source_hits VALUES (?, ?)", hit_batch)
                        connection.commit()
                        hit_batch.clear()
                    if source_files_scanned and source_files_scanned % 5000 == 0:
                        self.progress(
                            "scan", source_files_scanned, None, f"Scanned {source_files_scanned:,} files"
                        )

        if hit_batch:
            connection.executemany("INSERT OR IGNORE INTO source_hits VALUES (?, ?)", hit_batch)
            connection.commit()

        connection.execute(
            """
            UPDATE mapping_rows
            SET status = 'Not found', reason = ''
            WHERE status = 'Pending'
              AND norm_name NOT IN (SELECT norm_name FROM source_hits)
            """
        )
        connection.commit()

        def scalar(query: str) -> int:
            return int(connection.execute(query).fetchone()[0])

        summary = PreviewSummary(
            total_rows=scalar("SELECT COUNT(*) FROM mapping_rows"),
            valid_rows=scalar("SELECT COUNT(*) FROM mapping_rows WHERE status IN ('Pending', 'Not found')"),
            distinct_requested_files=scalar("SELECT COUNT(*) FROM target_names"),
            source_files_scanned=source_files_scanned,
            matched_rows=scalar("SELECT COUNT(*) FROM mapping_rows WHERE status = 'Pending'"),
            matched_source_files=scalar("SELECT COUNT(*) FROM source_hits"),
            not_found_rows=scalar("SELECT COUNT(*) FROM mapping_rows WHERE status = 'Not found'"),
            conflict_rows=scalar("SELECT COUNT(*) FROM mapping_rows WHERE status = 'Mapping conflict'"),
            invalid_rows=scalar("SELECT COUNT(*) FROM mapping_rows WHERE status = 'Move failed'"),
            folders_to_create=scalar(
                "SELECT COUNT(DISTINCT identifier) FROM mapping_rows WHERE status = 'Pending'"
            ),
            same_volume=os.path.splitdrive(str(cfg.source_folder.resolve()))[0].casefold()
            == os.path.splitdrive(str(cfg.target_folder.resolve()))[0].casefold(),
            elapsed_seconds=time.perf_counter() - started,
        )
        self.preview_summary = summary
        self.progress("preview_done", summary.total_rows, summary.total_rows, "Pre-check complete")
        return summary

    @staticmethod
    def _collision_safe_path(folder: Path, filename: str) -> Path:
        candidate = folder / filename
        if not candidate.exists():
            return candidate
        source = Path(filename)
        counter = 2
        while True:
            candidate = folder / f"{source.stem}_{counter}{source.suffix}"
            if not candidate.exists():
                return candidate
            counter += 1

    def _deliver(self, source: Path, folder: Path, move: bool) -> Path:
        """Reserve the destination exclusively; never overwrite a competing file."""
        folder.mkdir(parents=True, exist_ok=True)
        while True:
            destination = self._collision_safe_path(folder, source.name)
            if move:
                try:
                    os.link(source, destination)
                except FileExistsError:
                    continue
                except OSError:
                    pass  # Cross-volume or unsupported links: exclusive copy below.
                else:
                    source.unlink()
                    return destination
            try:
                output = destination.open("xb")
            except FileExistsError:
                continue
            try:
                with output, source.open("rb") as incoming:
                    while chunk := incoming.read(1024 * 1024):
                        self.control.checkpoint()
                        output.write(chunk)
                shutil.copystat(source, destination)
            except BaseException:
                destination.unlink(missing_ok=True)
                raise
            if move:
                source.unlink()
            return destination

    def execute(self) -> ExecutionSummary:
        if self.preview_summary is None:
            raise RuntimeError("Run pre-check before moving files")
        started = time.perf_counter()
        cfg = self.config
        connection = self._connect()
        cfg.target_folder.mkdir(parents=True, exist_ok=True)
        cfg.result_path.parent.mkdir(parents=True, exist_ok=True)
        log_path = cfg.result_path.with_name(f"{cfg.result_path.stem}_log.csv")
        total = int(
            connection.execute(
                """
                SELECT COUNT(*) FROM source_hits h
                WHERE EXISTS (
                    SELECT 1 FROM mapping_rows r
                    WHERE r.norm_name = h.norm_name AND r.status = 'Pending'
                )
                """
            ).fetchone()[0]
        )
        connection.executescript("""
            DROP TABLE IF EXISTS deliveries;
            CREATE TABLE deliveries AS
            SELECT DISTINCT h.norm_name, r.identifier, h.source_path,
                   'Pending' AS status, '' AS reason
            FROM source_hits h JOIN mapping_rows r ON r.norm_name = h.norm_name
            WHERE r.status = 'Pending';
            CREATE INDEX idx_deliveries_source ON deliveries(source_path);
            CREATE INDEX idx_deliveries_row ON deliveries(norm_name, identifier);
        """)
        groups = connection.execute(
            "SELECT norm_name, source_path, COUNT(*) FROM deliveries "
            "GROUP BY norm_name, source_path ORDER BY norm_name, source_path"
        )
        moved = 0
        cancelled = False

        write_header = not log_path.exists() or log_path.stat().st_size == 0
        with log_path.open("a", newline="", encoding="utf-8-sig") as log_file:
            writer = csv.writer(log_file)
            if write_header:
                writer.writerow(
                    ["Time", "Identifier", "File Name", "Source", "Destination", "Status", "Reason"]
                )

            for index, (norm_name, source_text, destination_count) in enumerate(groups, start=1):
                try:
                    self.control.checkpoint()
                except MappingCancelled:
                    cancelled = True
                    break
                source = Path(source_text)
                all_ok = True
                for (identifier,) in connection.execute(
                    "SELECT identifier FROM deliveries WHERE source_path = ?", (source_text,)
                ):
                    destination = cfg.target_folder / identifier / source.name
                    status, reason = "Mapped", ""
                    try:
                        self.control.checkpoint()
                        if not source.is_file():
                            raise FileNotFoundError("Source file is no longer available")
                        destination = self._deliver(source, destination.parent, destination_count == 1)
                    except MappingCancelled:
                        cancelled = True
                        all_ok = False
                        break
                    except Exception as exc:
                        status, reason = "Move failed", str(exc)
                        all_ok = False
                    connection.execute(
                        "UPDATE deliveries SET status = ?, reason = ? "
                        "WHERE source_path = ? AND identifier = ?",
                        (status, reason, source_text, identifier),
                    )
                    writer.writerow([datetime.now().isoformat(timespec="seconds"), identifier,
                                     source.name, str(source), str(destination), status, reason])
                    log_file.flush()
                if all_ok:
                    try:
                        if destination_count > 1:
                            source.unlink()
                        moved += 1
                    except OSError as exc:
                        connection.execute(
                            "UPDATE deliveries SET status = 'Move failed', reason = ? "
                            "WHERE source_path = ?", (f"Copies saved; source cleanup failed: {exc}", source_text)
                        )
                        writer.writerow([datetime.now().isoformat(timespec="seconds"), "", source.name,
                                         str(source), "", "Move failed", str(exc)])
                connection.commit()
                self.progress("move", index, total, f"Moved {moved:,} of {total:,} files")
                if cancelled:
                    break

        connection.execute("""
            UPDATE mapping_rows AS r SET
                status = CASE
                    WHEN EXISTS (SELECT 1 FROM deliveries d WHERE d.norm_name = r.norm_name
                         AND d.identifier = r.identifier AND d.status = 'Move failed') THEN 'Move failed'
                    WHEN EXISTS (SELECT 1 FROM deliveries d WHERE d.norm_name = r.norm_name
                         AND d.identifier = r.identifier AND d.status = 'Pending') THEN 'Pending'
                    ELSE 'Mapped' END,
                reason = COALESCE((SELECT group_concat(d.reason, '; ') FROM deliveries d
                    WHERE d.norm_name = r.norm_name AND d.identifier = r.identifier
                    AND d.reason != ''), '')
            WHERE r.status = 'Pending'
        """)
        connection.commit()

        self._write_result_workbook(connection)
        rows_mapped = int(
            connection.execute("SELECT COUNT(*) FROM mapping_rows WHERE status = 'Mapped'").fetchone()[0]
        )
        rows_failed = int(
            connection.execute(
                "SELECT COUNT(*) FROM mapping_rows "
                "WHERE status IN ('Move failed', 'Mapping conflict')"
            ).fetchone()[0]
        )
        rows_not_processed = int(
            connection.execute("SELECT COUNT(*) FROM mapping_rows WHERE status = 'Pending'").fetchone()[0]
        )
        summary = ExecutionSummary(
            files_moved=moved,
            rows_mapped=rows_mapped,
            rows_failed=rows_failed,
            rows_not_processed=rows_not_processed,
            result_path=cfg.result_path,
            log_path=log_path,
            cancelled=cancelled,
            elapsed_seconds=time.perf_counter() - started,
        )
        self.progress("done", total, total, "Processing complete")
        return summary

    def _write_result_workbook(self, connection: sqlite3.Connection) -> None:
        cfg = self.config
        self.progress("output", 0, None, "Writing result Excel")
        source_book = load_workbook(cfg.excel_path, read_only=True, data_only=False)
        output_book = Workbook(write_only=True)
        temporary_path = cfg.result_path.with_suffix(cfg.result_path.suffix + ".part")
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)

        try:
            for sheet_name in source_book.sheetnames:
                source_sheet = source_book[sheet_name]
                output_sheet = output_book.create_sheet(title=sheet_name)
                if sheet_name != cfg.sheet_name:
                    for row in source_sheet.iter_rows(values_only=True):
                        output_sheet.append(list(row))
                    continue

                status_cursor = iter(
                    connection.execute(
                        """
                        SELECT row_number,
                               CASE WHEN status = 'Pending' THEN 'Not processed' ELSE status END,
                               CASE WHEN status = 'Pending' THEN 'Cancelled before processing' ELSE reason END
                        FROM mapping_rows
                        ORDER BY row_number
                        """
                    )
                )
                next_status = next(status_cursor, None)
                first_row = next(source_sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
                header_values = list(first_row)
                status_header, reason_header = unique_output_headers(header_values)
                output_sheet.freeze_panes = "A2"
                for column_index, value in enumerate(
                    [*header_values, status_header, reason_header], start=1
                ):
                    text_length = len(_as_text(value))
                    if column_index == len(header_values) + 2:
                        width = 48
                    elif column_index == len(header_values) + 1:
                        width = 20
                    else:
                        width = min(40, max(14, text_length + 3))
                    output_sheet.column_dimensions[get_column_letter(column_index)].width = width

                styled: list[WriteOnlyCell] = []
                for value in [*header_values, status_header, reason_header]:
                    cell = WriteOnlyCell(output_sheet, value=value)
                    cell.fill = header_fill
                    cell.font = header_font
                    styled.append(cell)
                output_sheet.append(styled)

                for row_number, row in enumerate(
                    source_sheet.iter_rows(min_row=2, values_only=True), start=2
                ):
                    values = list(row)
                    status = ""
                    reason = ""
                    if next_status is not None and next_status[0] == row_number:
                        status = next_status[1]
                        reason = next_status[2]
                        next_status = next(status_cursor, None)
                    output_sheet.append([*values, status, reason])
                    if row_number % 5000 == 0:
                        output_total = (
                            source_sheet.max_row - 1
                            if isinstance(source_sheet.max_row, int)
                            else None
                        )
                        self.progress(
                            "output", row_number - 1, output_total, f"Wrote {row_number - 1:,} rows"
                        )

            output_book.save(temporary_path)
            os.replace(temporary_path, cfg.result_path)
        finally:
            source_book.close()
            if temporary_path.exists():
                temporary_path.unlink(missing_ok=True)


def workbook_metadata(path: Path) -> dict[str, list[tuple[int, str]]]:
    """Read worksheet names and first-row headers without loading the workbook into memory."""
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        result: dict[str, list[tuple[int, str]]] = {}
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            result[sheet_name] = [
                (index, _as_text(value) or f"Column {index}")
                for index, value in enumerate(first_row, start=1)
            ]
        return result
    finally:
        workbook.close()
