from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from mapping_core import MappingConfig, MappingEngine, MappingCancelled


class MappingEngineTest(unittest.TestCase):
    def test_exact_names_multiple_wo_and_all_sources(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source, target = root / 'source', root / 'target'
            for sub, content in [('a', 'first'), ('b', 'second')]:
                (source / sub).mkdir(parents=True)
                (source / sub / 'Report.pdf').write_text(content)
            (source / 'c').mkdir()
            (source / 'c' / 'report.PDF').write_text('different case')
            (target / 'WO1').mkdir(parents=True)
            (target / 'WO1' / 'REPORT.PDF').write_text('existing')
            book = Workbook()
            book.active.title = 'Data'
            for row in [('WO', 'File'), ('WO1', 'Report.pdf'), ('WO2', 'Report.pdf'),
                        ('WO1', 'Report.pdf'), ('WO3', 'report.PDF'), ('WO4', 'report.pdf')]:
                book.active.append(row)
            book.save(root / 'input.xlsx')
            with MappingEngine(MappingConfig(root / 'input.xlsx', 'Data', 1, 2,
                    source, target, root / 'result.xlsx', include_subfolders=True)) as engine:
                preview = engine.precheck()
                self.assertEqual((preview.matched_rows, preview.not_found_rows, preview.conflict_rows), (4, 1, 0))
                result = engine.execute()
                self.assertEqual((result.files_moved, result.rows_mapped), (3, 4))
                self.assertEqual((target / 'WO1' / 'REPORT.PDF').read_text(), 'existing')
                self.assertEqual({p.read_text() for p in (target / 'WO1').iterdir()}, {'existing', 'first', 'second'})
                self.assertEqual({p.read_text() for p in (target / 'WO2').iterdir()}, {'first', 'second'})
                self.assertEqual((target / 'WO3' / 'report.PDF').read_text(), 'different case')
                self.assertFalse((target / 'WO4').exists())
                self.assertFalse(list(source.rglob('*.pdf')))

    def test_failed_delivery_keeps_source_and_successful_wo(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / 'source'
            source.mkdir()
            (source / 'A.txt').write_text('original')
            book = Workbook()
            book.active.title = 'Data'
            for row in [('WO', 'File'), ('WO1', 'A.txt'), ('WO2', 'A.txt')]:
                book.active.append(row)
            book.save(root / 'input.xlsx')
            with MappingEngine(MappingConfig(root / 'input.xlsx', 'Data', 1, 2,
                    source, root / 'target', root / 'result.xlsx')) as engine:
                engine.precheck()
                deliver = engine._deliver
                def fail_one(src, folder, move):
                    if folder.name == 'WO2':
                        raise PermissionError('test failure')
                    return deliver(src, folder, move)
                with patch.object(engine, '_deliver', side_effect=fail_one):
                    result = engine.execute()
                self.assertEqual((result.rows_mapped, result.rows_failed), (1, 1))
                self.assertTrue((source / 'A.txt').exists())
                self.assertEqual((root / 'target/WO1/A.txt').read_text(), 'original')

    def test_cancel_during_shared_file_retains_source(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / 'source'
            source.mkdir()
            (source / 'A.txt').write_text('original')
            book = Workbook()
            book.active.title = 'Data'
            for row in [('WO', 'File'), ('WO1', 'A.txt'), ('WO2', 'A.txt')]:
                book.active.append(row)
            book.save(root / 'input.xlsx')
            with MappingEngine(MappingConfig(root / 'input.xlsx', 'Data', 1, 2,
                    source, root / 'target', root / 'result.xlsx')) as engine:
                engine.precheck()
                deliver = engine._deliver
                def cancel_second(src, folder, move):
                    if folder.name == 'WO2':
                        raise MappingCancelled()
                    return deliver(src, folder, move)
                with patch.object(engine, '_deliver', side_effect=cancel_second):
                    result = engine.execute()
                self.assertTrue(result.cancelled)
                self.assertEqual((result.rows_mapped, result.rows_not_processed), (1, 1))
                self.assertTrue((source / 'A.txt').exists())

    def test_custom_columns_and_recursive_scan(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            workbook_path = root / "mapping.xlsx"
            source = root / "source"
            nested = source / "nested"
            target = root / "target"
            nested.mkdir(parents=True)
            target.mkdir()

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Data"
            sheet.append(["Notes", "Exact File Name", "Unused", "Report Number"])
            sheet.append(["test", "Report001.txt", "x", "REPORT-31000001"])
            workbook.save(workbook_path)
            (nested / "Report001.txt").write_text("test", encoding="utf-8")

            config = MappingConfig(
                excel_path=workbook_path,
                sheet_name="Data",
                identifier_column=4,
                filename_column=2,
                source_folder=source,
                target_folder=target,
                result_path=root / "result.xlsx",
                include_subfolders=True,
            )
            engine = MappingEngine(config)
            try:
                preview = engine.precheck()
                self.assertEqual(preview.total_rows, 1)
                self.assertEqual(preview.source_files_scanned, 1)
                self.assertEqual(preview.matched_source_files, 1)

                summary = engine.execute()
                self.assertFalse(summary.cancelled)
                self.assertEqual(summary.files_moved, 1)
                self.assertTrue((target / "REPORT-31000001" / "Report001.txt").is_file())
                self.assertTrue(summary.result_path.is_file())
                self.assertTrue(summary.log_path.is_file())
            finally:
                engine.close()


if __name__ == "__main__":
    unittest.main()
